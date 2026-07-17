"""Parser tests: bytes -> ParsedDoc with cell/page lineage, magic-byte routing."""

import io

from cre_mcp.truth.parsers import parse_bytes, sniff_ext


def _xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "RR"
    ws.append(["Tenant", "Base Rent"])
    ws.append(["Starbucks", 72000])
    ws.append(["Chase Bank", 140000])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_sniff_ext_uses_magic_not_extension():
    assert sniff_ext(b"%PDF-1.7\n...", hint=".xlsx") == "pdf"  # magic wins over lying hint
    assert sniff_ext(_xlsx_bytes(), hint=".txt") == "xlsx"
    assert sniff_ext(b"a,b,c\n1,2,3", hint=".csv") == "csv"


def test_csv_parse_keeps_cell_refs():
    data = b"Line Item,Amount\nGross Rent,252000\nManagement Fee,-12600\n"
    doc = parse_bytes(data, ext_hint=".csv")
    assert doc.ext == "csv"
    assert doc.n_pages == 1
    table = doc.tables[0]
    # header row + 2 data rows
    assert len(table.rows) == 3
    amount_cell = table.rows[1][1]  # "252000"
    assert amount_cell.ref == "B2"
    assert amount_cell.value == "252000"


def test_xlsx_parse_openpyxl_cell_lineage():
    doc = parse_bytes(_xlsx_bytes(), ext_hint=".xlsx")
    assert doc.ext == "xlsx"
    table = doc.tables[0]
    rent_cell = table.rows[1][1]  # Starbucks base rent
    assert rent_cell.ref == "RR!B2"
    assert rent_cell.value == "72000"
    assert "Starbucks" in doc.full_text
