"""Real-workbook forensics and published-table cost-seg tests."""

from openpyxl import Workbook
import pytest

from cre_mcp.analytics.costseg import cost_seg_preview
from cre_mcp.analytics.model_forensics import audit_seller_model


def test_audit_seller_model_finds_planted_real_xlsx_defects_with_citations(tmp_path):
    path = tmp_path / "planted_seller_model.xlsx"
    workbook = Workbook()
    model = workbook.active
    model.title = "Seller Model"
    model["B2"] = "=A2*2"
    model["B3"] = 777  # planted overwrite in the vertical formula run
    model["B4"] = "=A4*2"
    model["D2"] = "='[OldUnderwriting.xlsx]Rent Roll'!A1"
    model["F2"] = "=F3+1"
    model["F3"] = "=F2+1"
    model.row_dimensions[6].hidden = True
    model.column_dimensions["C"].hidden = True

    hidden = workbook.create_sheet("Hidden Calc")
    hidden.sheet_state = "hidden"
    very_hidden = workbook.create_sheet("Very Hidden")
    very_hidden.sheet_state = "veryHidden"
    workbook.save(path)

    result = audit_seller_model(path)
    findings = result["findings"]
    by_type = {
        kind: [item for item in findings if item["type"] == kind]
        for kind in {item["type"] for item in findings}
    }

    assert result["status"] == "FINDINGS"
    assert by_type["hidden_row"][0]["cell_citation"] == "'Seller Model'!A6"
    assert by_type["hidden_column"][0]["cell_citation"] == "'Seller Model'!C1"
    assert by_type["hardcoded_in_formula_range"][0]["cell_citation"] == "'Seller Model'!B3"
    assert by_type["hidden_sheet"][0]["cell_citation"] == "'Hidden Calc'!A1"
    assert by_type["very_hidden_sheet"][0]["cell_citation"] == "'Very Hidden'!A1"
    assert by_type["external_link"][0]["cell_citation"] == "'Seller Model'!D2"
    circular = by_type["circular_reference"][0]
    assert set(circular["related_cells"]) == {
        "'Seller Model'!F2",
        "'Seller Model'!F3",
    }
    assert all(item["cell_citation"] for item in findings)


def test_cost_seg_five_year_schedule_matches_published_macrs_table_year_1_to_n():
    result = cost_seg_preview(
        purchase_price=1_000_000,
        land_pct=20,
        components=[{"class": "5yr", "pct": 100}],
    )

    expected_rates = [0.20, 0.32, 0.192, 0.1152, 0.1152, 0.0576]
    expected_dollars = [160_000, 256_000, 153_600, 92_160, 92_160, 46_080]
    component = result["components"][0]

    assert component["macrs_rates"] == expected_rates
    assert [row["total_depreciation"] for row in component["schedule"]] == expected_dollars
    assert sum(row["total_depreciation"] for row in component["schedule"]) == 800_000
    assert result["yearly_schedule"][-1]["remaining_depreciable_basis"] == 0
    assert "Table A-1" in component["citation"]
    assert result["filing_requirement"] == (
        "engineering-based cost-seg study + CPA required for filing"
    )
    assert result["after_tax_hook"]["module"] == "cre_mcp.ops.after_tax_returns"


def test_cost_seg_bonus_is_explicit_and_unknown_component_keys_are_surfaced():
    result = cost_seg_preview(
        100_000,
        0,
        [{"class": "5yr", "pct": 1, "seller_guess": "yes"}],
        bonus_rate=80,
        apply_bonus=True,
    )

    first_year = result["components"][0]["schedule"][0]
    # 80,000 bonus + 20% MACRS on the remaining 20,000.
    assert first_year["bonus_depreciation"] == 80_000
    assert first_year["regular_macrs_depreciation"] == 4_000
    assert result["unrecognized_inputs"] == ["components[0].seller_guess"]
    assert any("law changes" in warning for warning in result["warnings"])


@pytest.mark.parametrize(
    "args, expected",
    [
        ((None, 20, [{"class": "5yr", "pct": 100}]), "purchase_price"),
        ((100_000, 20, None), "components or asset_type"),
        ((100_000, 20, [{"class": "5yr", "pct": 50}]), "sum to 100%"),
    ],
)
def test_cost_seg_nullable_and_invalid_inputs_stay_inside_error_boundary(args, expected):
    result = cost_seg_preview(*args)
    assert set(result) == {"error"}
    assert expected in result["error"]


def test_cost_seg_profiles_are_exposed_and_39_year_is_not_mislabeled_half_year():
    result = cost_seg_preview(1_000_000, 20, asset_type="retail")

    assert result["input_mode"] == "screening_profile"
    assert "retail" in result["available_asset_type_profiles"]
    real_property = next(item for item in result["components"] if item["class"] == "39yr")
    assert "mid-month" in real_property["method"]
    assert real_property["bonus_rate_applied"] == 0
