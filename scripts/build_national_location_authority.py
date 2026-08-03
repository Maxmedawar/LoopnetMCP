#!/usr/bin/env python3
"""Build deterministic national county, place, ZIP, tract, and CBSA relationships.

The runtime artifact contains five pipe-delimited row types:

* ``C|county_fips|state|county_name``
* ``P|state|place_name|place_geoid|county_fips``
* ``Z|zip5|state|county_fips``
* ``T|tract_geoid|state|county_fips|zip5_or_empty``
* ``B|county_fips|cbsa|``

Only relationships backed by the supplied Census Bureau source files are
emitted. ZIP rows are additionally restricted to ZIPs already assigned in the
checked-in postal authority, so a 2020 ZCTA cannot create a postal assignment.
"""

from __future__ import annotations

import argparse
import csv
import gzip
import hashlib
from pathlib import Path

from openpyxl import load_workbook


_SOURCE_SHA256 = {
    "counties": "9f6e5f6eb6ac2f5e9a36d5fd01dec77991bddc75118f748a069441a4782970d6",
    "places": "9996494d33cf6ee4527491508aa1b8dae85dcff3a979886e19b0b375e5d06ec6",
    "zcta_counties": "3ed41278d637dc249e0323306f68be8a6c234e3090f4de88ef328dee71aeaaaf",
    "zcta_tracts": "6a25d8c3fff4cf612c4d2dccc2c0cd6cb5bc99b807ff3d5d107a2e9b9d68dde0",
    "cbsa": "952c4b1e78acbb54e6ec9412434b7602fedacbf021736351a63c181bdb753629",
}
_ALLOWED_STATES = frozenset(
    {
        "AL", "AK", "AZ", "AR", "CA", "CO", "CT", "DE", "DC", "FL",
        "GA", "HI", "ID", "IL", "IN", "IA", "KS", "KY", "LA", "ME",
        "MD", "MA", "MI", "MN", "MS", "MO", "MT", "NE", "NV", "NH",
        "NJ", "NM", "NY", "NC", "ND", "OH", "OK", "OR", "PA", "RI",
        "SC", "SD", "TN", "TX", "UT", "VT", "VA", "WA", "WV", "WI",
        "WY", "PR",
    }
)
_PLACE_SUFFIXES = tuple(
    sorted(
        (
            " consolidated government",
            " metropolitan government",
            " unified government",
            " city and borough",
            " metro government",
            " metro township",
            " urban county",
            " consolidated municipality",
            " municipality",
            " corporation",
            " zona urbana",
            " comunidad",
            " borough",
            " village",
            " town",
            " city",
            " cdp",
        ),
        key=len,
        reverse=True,
    )
)
_COUNTY_SUFFIXES = tuple(
    sorted(
        (
            " city and borough",
            " census area",
            " municipality",
            " municipio",
            " borough",
            " parish",
            " county",
            " city",
        ),
        key=len,
        reverse=True,
    )
)


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _check_source(name: str, path: Path) -> None:
    actual = _sha256(path)
    expected = _SOURCE_SHA256[name]
    if actual != expected:
        raise ValueError(
            f"unexpected {name} source hash: expected {expected}, got {actual}"
        )


def _safe_text(value: object, *, field: str) -> str:
    if not isinstance(value, str):
        raise ValueError(f"missing {field}")
    normalized = " ".join(value.split())
    if not normalized or any(character in normalized for character in "|\r\n"):
        raise ValueError(f"invalid {field}: {value!r}")
    return normalized


def _bare_place_name(value: object) -> str:
    place = _safe_text(value, field="place name")
    if place.casefold().endswith(" (balance)"):
        place = place[: -len(" (balance)")].rstrip()
    folded = place.casefold()
    for suffix in _PLACE_SUFFIXES:
        if folded.endswith(suffix):
            place = place[: -len(suffix)].rstrip()
            break
    return _safe_text(place, field="bare place name")


def _bare_county_name(value: object) -> str:
    county = _safe_text(value, field="county name")
    if county == "District of Columbia":
        return county
    folded = county.casefold()
    for suffix in _COUNTY_SUFFIXES:
        if folded.endswith(suffix):
            return county
    raise ValueError(f"unsupported county-equivalent name: {county!r}")


def _postal_states(path: Path) -> dict[str, str]:
    states: dict[str, str] = {}
    decoded = gzip.decompress(path.read_bytes()).decode("utf-8")
    for line in decoded.splitlines():
        zip5, state, _city = line.split("|", 2)
        prior = states.setdefault(zip5, state)
        if prior != state:
            raise ValueError(f"postal ZIP {zip5} crosses configured states")
    return states


def _county_rows(path: Path) -> tuple[dict[str, tuple[str, str]], set[str]]:
    counties: dict[str, tuple[str, str]] = {}
    rows: set[str] = set()
    with path.open(encoding="utf-8-sig", newline="") as source:
        for record in csv.DictReader(source, delimiter="|"):
            state = _safe_text(record.get("STATE"), field="county state")
            if state not in _ALLOWED_STATES:
                continue
            fips = f"{record['STATEFP']}{record['COUNTYFP']}"
            if len(fips) != 5 or not fips.isascii() or not fips.isdecimal():
                raise ValueError(f"invalid county FIPS: {fips!r}")
            county = _bare_county_name(record.get("COUNTYNAME"))
            identity = (state, county)
            prior = counties.setdefault(fips, identity)
            if prior != identity:
                raise ValueError(f"conflicting county identity for {fips}")
            rows.add(f"C|{fips}|{state}|{county}")
    return counties, rows


def _place_rows(
    path: Path,
    counties: dict[str, tuple[str, str]],
) -> set[str]:
    rows: set[str] = set()
    with path.open(encoding="utf-8-sig", newline="") as source:
        for record in csv.DictReader(source, delimiter="|"):
            state = _safe_text(record.get("STATE"), field="place state")
            if state not in _ALLOWED_STATES or record.get("FUNCSTAT") not in {"A", "S"}:
                continue
            fips = f"{record['STATEFP']}{record['COUNTYFP']}"
            if counties.get(fips, (None,))[0] != state:
                raise ValueError(f"place references unknown county {fips}")
            place = _bare_place_name(record.get("PLACENAME"))
            place_geoid = f"{record['STATEFP']}{record['PLACEFP']}"
            if (
                len(place_geoid) != 7
                or not place_geoid.isascii()
                or not place_geoid.isdecimal()
                or place_geoid[:2] != fips[:2]
            ):
                raise ValueError(f"invalid place GEOID: {place_geoid!r}")
            rows.add(f"P|{state}|{place}|{place_geoid}|{fips}")
    return rows


def _zcta_rows(
    path: Path,
    counties: dict[str, tuple[str, str]],
    postal_states: dict[str, str],
) -> set[str]:
    rows: set[str] = set()
    with path.open(encoding="utf-8-sig", newline="") as source:
        for record in csv.DictReader(source, delimiter="|"):
            zip5 = record.get("GEOID_ZCTA5_20") or ""
            fips = record.get("GEOID_COUNTY_20") or ""
            if len(zip5) != 5 or not zip5.isascii() or not zip5.isdecimal():
                continue
            county = counties.get(fips)
            state = postal_states.get(zip5)
            if county is None or state is None or county[0] != state:
                continue
            try:
                land_area = int(record["AREALAND_PART"])
            except (KeyError, TypeError, ValueError):
                raise ValueError("invalid ZCTA-to-county land area") from None
            if land_area > 0:
                rows.add(f"Z|{zip5}|{state}|{fips}")
    return rows


def _cbsa_rows(
    path: Path,
    counties: dict[str, tuple[str, str]],
) -> set[str]:
    rows: set[str] = set()
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    county_to_cbsa: dict[str, str] = {}
    for record in sheet.iter_rows(min_row=4, values_only=True):
        cbsa_raw, state_fips_raw, county_fips_raw = record[0], record[9], record[10]
        if cbsa_raw is None or state_fips_raw is None or county_fips_raw is None:
            continue
        cbsa = str(cbsa_raw).zfill(5)
        fips = f"{str(state_fips_raw).zfill(2)}{str(county_fips_raw).zfill(3)}"
        if counties.get(fips) is None:
            # July 2023 Connecticut planning regions do not share the 2020
            # county authority used by the other relationship sources.
            continue
        if len(cbsa) != 5 or not cbsa.isascii() or not cbsa.isdecimal():
            raise ValueError(f"invalid CBSA: {cbsa!r}")
        prior = county_to_cbsa.setdefault(fips, cbsa)
        if prior != cbsa:
            raise ValueError(f"county {fips} belongs to multiple CBSAs")
        rows.add(f"B|{fips}|{cbsa}|")
    workbook.close()
    return rows


def _tract_rows(
    path: Path,
    counties: dict[str, tuple[str, str]],
    postal_states: dict[str, str],
) -> set[str]:
    """Build exact existing-tract and positive-land ZCTA relations."""

    rows: set[str] = set()
    with path.open(encoding="utf-8-sig", newline="") as source:
        for record in csv.DictReader(source, delimiter="|"):
            tract = record.get("GEOID_TRACT_20") or ""
            if not tract:
                continue
            if len(tract) != 11 or not tract.isascii() or not tract.isdecimal():
                raise ValueError(f"invalid tract GEOID: {tract!r}")
            county_fips = tract[:5]
            county = counties.get(county_fips)
            if county is None:
                continue
            try:
                land_area = int(record["AREALAND_PART"])
            except (KeyError, TypeError, ValueError):
                raise ValueError("invalid ZCTA-to-tract land area") from None
            if land_area <= 0:
                continue
            zip5 = record.get("GEOID_ZCTA5_20") or ""
            if zip5 and (
                len(zip5) != 5
                or not zip5.isascii()
                or not zip5.isdecimal()
                or postal_states.get(zip5) != county[0]
            ):
                continue
            rows.add(f"T|{tract}|{county[0]}|{county_fips}|{zip5}")
    return rows


def build(
    *,
    counties_path: Path,
    places_path: Path,
    zcta_counties_path: Path,
    zcta_tracts_path: Path,
    cbsa_path: Path,
    postal_path: Path,
    output_path: Path,
) -> None:
    _check_source("counties", counties_path)
    _check_source("places", places_path)
    _check_source("zcta_counties", zcta_counties_path)
    _check_source("zcta_tracts", zcta_tracts_path)
    _check_source("cbsa", cbsa_path)

    counties, rows = _county_rows(counties_path)
    postal_states = _postal_states(postal_path)
    rows.update(_place_rows(places_path, counties))
    rows.update(
        _zcta_rows(
            zcta_counties_path,
            counties,
            postal_states,
        )
    )
    rows.update(_tract_rows(zcta_tracts_path, counties, postal_states))
    rows.update(_cbsa_rows(cbsa_path, counties))

    payload = ("\n".join(sorted(rows)) + "\n").encode("utf-8")
    output_path.write_bytes(gzip.compress(payload, compresslevel=9, mtime=0))
    counts = {
        prefix: sum(row.startswith(f"{prefix}|") for row in rows)
        for prefix in ("C", "P", "Z", "T", "B")
    }
    print(
        f"wrote {len(rows)} rows {counts} to {output_path} "
        f"sha256={_sha256(output_path)}"
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--counties", required=True, type=Path)
    parser.add_argument("--places", required=True, type=Path)
    parser.add_argument("--zcta-counties", required=True, type=Path)
    parser.add_argument("--zcta-tracts", required=True, type=Path)
    parser.add_argument("--cbsa", required=True, type=Path)
    parser.add_argument("--postal", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    args = parser.parse_args()
    build(
        counties_path=args.counties,
        places_path=args.places,
        zcta_counties_path=args.zcta_counties,
        zcta_tracts_path=args.zcta_tracts,
        cbsa_path=args.cbsa,
        postal_path=args.postal,
        output_path=args.output,
    )


if __name__ == "__main__":
    main()
