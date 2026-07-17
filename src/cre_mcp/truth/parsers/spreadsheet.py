"""Spreadsheet + CSV parsing with cell-level lineage.

Primary path: openpyxl (read-only, cached values). Fallback: stdlib zipfile +
xml.etree over the OOXML parts, so xlsx still parses if the ``truth`` extra is
absent. CSV uses the stdlib ``csv`` module (zero dependency).
"""

from __future__ import annotations

import csv
import io
import re
import zipfile
from xml.etree import ElementTree as ET

from cre_mcp.truth.parsers.base import Cell, Page, ParsedDoc, Table


def _col_letter(index: int) -> str:
    """1-based column index -> spreadsheet letter (1 -> A, 27 -> AA)."""
    letters = ""
    while index > 0:
        index, rem = divmod(index - 1, 26)
        letters = chr(65 + rem) + letters
    return letters


def _rows_to_table(
    rows: list[list[str]], *, sheet: str | None, page: int | None = None
) -> Table:
    table = Table(sheet=sheet, page=page)
    for r_idx, row in enumerate(rows, start=1):
        cells: list[Cell] = []
        for c_idx, value in enumerate(row, start=1):
            ref = f"{_col_letter(c_idx)}{r_idx}"
            full_ref = f"{sheet}!{ref}" if sheet else ref
            cells.append(
                Cell(ref=full_ref, value=str(value), row=r_idx, col=c_idx, sheet=sheet, page=page)
            )
        table.rows.append(cells)
    return table


def _table_to_text(rows: list[list[str]]) -> str:
    return "\n".join("\t".join(str(v) for v in row) for row in rows)


def parse_csv(data: bytes, sha256: str) -> ParsedDoc:
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    rows = [list(row) for row in reader]
    table = _rows_to_table(rows, sheet=None)
    page = Page(number=1, text=_table_to_text(rows))
    return ParsedDoc(
        sha256=sha256, ext="csv", pages=[page], tables=[table], n_pages=1
    )


def parse_xlsx(data: bytes, sha256: str) -> ParsedDoc:
    try:
        return _parse_xlsx_openpyxl(data, sha256)
    except ImportError:
        return _parse_xlsx_stdlib(data, sha256)


def _parse_xlsx_openpyxl(data: bytes, sha256: str) -> ParsedDoc:
    from openpyxl import load_workbook  # lazy: part of the optional truth extra

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    tables: list[Table] = []
    pages: list[Page] = []
    for s_idx, ws in enumerate(wb.worksheets, start=1):
        rows: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            rows.append(["" if v is None else v for v in row])
        tables.append(_rows_to_table(rows, sheet=ws.title))
        pages.append(Page(number=s_idx, text=_table_to_text(rows)))
    wb.close()
    if not pages:
        pages = [Page(number=1, text="")]
    return ParsedDoc(
        sha256=sha256, ext="xlsx", pages=pages, tables=tables, n_pages=len(pages)
    )


_NS = {"m": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
_CELL_REF = re.compile(r"([A-Z]+)(\d+)")


def _ref_to_rc(ref: str) -> tuple[int, int]:
    match = _CELL_REF.match(ref)
    if not match:
        return (1, 1)
    col_s, row_s = match.groups()
    col = 0
    for ch in col_s:
        col = col * 26 + (ord(ch) - 64)
    return (int(row_s), col)


def _parse_xlsx_stdlib(data: bytes, sha256: str) -> ParsedDoc:
    """Dependency-free xlsx read: unzip + parse sharedStrings & the first sheet."""
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        shared: list[str] = []
        if "xl/sharedStrings.xml" in zf.namelist():
            sroot = ET.fromstring(zf.read("xl/sharedStrings.xml"))
            for si in sroot.findall("m:si", _NS):
                shared.append("".join(t.text or "" for t in si.iter() if t.tag.endswith("}t")))
        sheet_names = [n for n in zf.namelist() if re.match(r"xl/worksheets/sheet\d+\.xml", n)]
        sheet_names.sort()
        if not sheet_names:
            return ParsedDoc(sha256=sha256, ext="xlsx", pages=[Page(1, "")], n_pages=1,
                             parse_status="failed", warnings=["no worksheet parts"])
        root = ET.fromstring(zf.read(sheet_names[0]))
        grid: dict[int, dict[int, str]] = {}
        max_col = 0
        for c in root.iter("{%s}c" % _NS["m"]):
            ref = c.get("r", "A1")
            row, col = _ref_to_rc(ref)
            max_col = max(max_col, col)
            v = c.find("m:v", _NS)
            if v is None or v.text is None:
                continue
            if c.get("t") == "s":
                idx = int(v.text)
                value = shared[idx] if 0 <= idx < len(shared) else ""
            else:
                value = v.text
            grid.setdefault(row, {})[col] = value
    rows: list[list[str]] = []
    for r in range(1, (max(grid) if grid else 0) + 1):
        rows.append([grid.get(r, {}).get(c, "") for c in range(1, max_col + 1)])
    table = _rows_to_table(rows, sheet="Sheet1")
    return ParsedDoc(
        sha256=sha256, ext="xlsx", pages=[Page(1, _table_to_text(rows))],
        tables=[table], n_pages=1,
    )


__all__ = ["parse_csv", "parse_xlsx"]
