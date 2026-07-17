"""Read-only structural forensics for seller-supplied Excel workbooks.

The audit deliberately inspects formulas rather than calculated values.  It is
therefore useful as a triage pass, not as a substitute for tracing the model in
Excel or reconciling it to source documents.
"""

from __future__ import annotations

from collections.abc import Iterable
from numbers import Number
from pathlib import Path
import re
from typing import Any

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils import get_column_letter, range_boundaries


_CELL_OR_RANGE = re.compile(
    r"^(?:(?P<sheet>'(?:[^']|'')+'|[^!]+)!)?"
    r"(?P<start>\$?[A-Za-z]{1,3}\$?[1-9][0-9]*)"
    r"(?::(?P<end>\$?[A-Za-z]{1,3}\$?[1-9][0-9]*))?$"
)
_EXTERNAL_BOOK = re.compile(r"\[[^\]]+\]")


def _citation(sheet: str, coordinate: str) -> str:
    """Return an Excel-style cell citation, quoting sheet names when needed."""

    if re.fullmatch(r"[A-Za-z_][A-Za-z0-9_.]*", sheet):
        display = sheet
    else:
        display = "'" + sheet.replace("'", "''") + "'"
    return f"{display}!{coordinate}"


def _finding(
    kind: str,
    sheet: str,
    coordinate: str,
    detail: str,
    *,
    severity: str = "review",
    related_cells: Iterable[str] = (),
) -> dict[str, Any]:
    return {
        "type": kind,
        "category": kind,
        "severity": severity,
        "cell_citation": _citation(sheet, coordinate),
        "related_cells": list(related_cells),
        "detail": detail,
    }


def _is_formula(value: Any) -> bool:
    return isinstance(value, str) and value.startswith("=")


def _is_hardcoded_value(value: Any) -> bool:
    """Identify plausible formula replacements while excluding empty cells."""

    if value is None or _is_formula(value):
        return False
    # Numeric constants are the usual defect.  Strings and dates are retained
    # too because a seller model can hardcode labels such as "N/A" into a
    # formula run; the opposite-neighbor test prevents ordinary labels from
    # being reported in most layouts.
    return not (isinstance(value, str) and not value.strip())


def _opposite_formula_neighbors(worksheet: Any, row: int, column: int) -> list[str]:
    pairs = (
        ((row, column - 1), (row, column + 1)),
        ((row - 1, column), (row + 1, column)),
    )
    for first, second in pairs:
        if min(first + second) < 1:
            continue
        first_cell = worksheet.cell(*first)
        second_cell = worksheet.cell(*second)
        if _is_formula(first_cell.value) and _is_formula(second_cell.value):
            return [
                _citation(worksheet.title, first_cell.coordinate),
                _citation(worksheet.title, second_cell.coordinate),
            ]
    return []


def _formula_references(
    formula: str,
    current_sheet: str,
    sheet_names: set[str],
) -> set[str]:
    """Extract A1 references textually, including bounded rectangular ranges."""

    references: set[str] = set()
    try:
        tokens = Tokenizer(formula).items
    except Exception:
        return references

    for token in tokens:
        if token.type != "OPERAND" or token.subtype != "RANGE":
            continue
        value = token.value.strip()
        if _EXTERNAL_BOOK.search(value):
            continue
        match = _CELL_OR_RANGE.fullmatch(value)
        if not match:
            continue
        raw_sheet = match.group("sheet")
        if raw_sheet is None:
            sheet = current_sheet
        elif raw_sheet.startswith("'") and raw_sheet.endswith("'"):
            sheet = raw_sheet[1:-1].replace("''", "'")
        else:
            sheet = raw_sheet
        if sheet not in sheet_names:
            continue

        start = match.group("start").replace("$", "")
        end = (match.group("end") or start).replace("$", "")
        min_col, min_row, max_col, max_row = range_boundaries(f"{start}:{end}")
        # A whole-column formula range can be enormous.  The cap keeps this
        # structural pass bounded; endpoints are still checked when capped.
        count = (max_col - min_col + 1) * (max_row - min_row + 1)
        if count > 10_000:
            coordinates = (start, end)
        else:
            coordinates = (
                f"{get_column_letter(column)}{row}"
                for row in range(min_row, max_row + 1)
                for column in range(min_col, max_col + 1)
            )
        references.update(_citation(sheet, coordinate) for coordinate in coordinates)
    return references


def _strongly_connected_components(graph: dict[str, set[str]]) -> list[list[str]]:
    """Tarjan SCCs for textual formula-reference cycles."""

    index = 0
    stack: list[str] = []
    on_stack: set[str] = set()
    indexes: dict[str, int] = {}
    lowlinks: dict[str, int] = {}
    components: list[list[str]] = []

    def visit(node: str) -> None:
        nonlocal index
        indexes[node] = index
        lowlinks[node] = index
        index += 1
        stack.append(node)
        on_stack.add(node)

        for neighbor in graph.get(node, set()):
            if neighbor not in graph:
                continue
            if neighbor not in indexes:
                visit(neighbor)
                lowlinks[node] = min(lowlinks[node], lowlinks[neighbor])
            elif neighbor in on_stack:
                lowlinks[node] = min(lowlinks[node], indexes[neighbor])

        if lowlinks[node] != indexes[node]:
            return
        component: list[str] = []
        while stack:
            member = stack.pop()
            on_stack.remove(member)
            component.append(member)
            if member == node:
                break
        components.append(component)

    for node in graph:
        if node not in indexes:
            visit(node)
    return components


def _audit_seller_model(xlsx_path: str | Path | None) -> dict[str, Any]:
    if xlsx_path is None or isinstance(xlsx_path, bool):
        raise ValueError("xlsx_path is required")
    path = Path(xlsx_path).expanduser()
    if not path.is_file():
        raise ValueError(f"xlsx_path does not exist or is not a file: {path}")
    if path.suffix.lower() not in {".xlsx", ".xlsm", ".xltx", ".xltm"}:
        raise ValueError("xlsx_path must be an openpyxl-compatible Excel workbook")

    workbook = load_workbook(
        path,
        data_only=False,
        read_only=False,
        keep_links=True,
        keep_vba=path.suffix.lower() in {".xlsm", ".xltm"},
    )
    findings: list[dict[str, Any]] = []
    formula_graph: dict[str, set[str]] = {}
    external_formula_cells: set[str] = set()
    sheet_names = set(workbook.sheetnames)

    for worksheet in workbook.worksheets:
        if worksheet.sheet_state == "hidden":
            findings.append(
                _finding(
                    "hidden_sheet",
                    worksheet.title,
                    "A1",
                    f"Worksheet {worksheet.title!r} is hidden.",
                )
            )
        elif worksheet.sheet_state == "veryHidden":
            findings.append(
                _finding(
                    "very_hidden_sheet",
                    worksheet.title,
                    "A1",
                    (
                        f"Worksheet {worksheet.title!r} is very-hidden and cannot be "
                        "made visible through Excel's ordinary sheet menu."
                    ),
                    severity="high",
                )
            )

        for row_number, dimension in worksheet.row_dimensions.items():
            if dimension.hidden:
                findings.append(
                    _finding(
                        "hidden_row",
                        worksheet.title,
                        f"A{row_number}",
                        f"Row {row_number} is hidden.",
                    )
                )

        for column_key, dimension in worksheet.column_dimensions.items():
            if not dimension.hidden:
                continue
            minimum = dimension.min or 1
            maximum = dimension.max or minimum
            for column_number in range(minimum, maximum + 1):
                column_letter = get_column_letter(column_number)
                findings.append(
                    _finding(
                        "hidden_column",
                        worksheet.title,
                        f"{column_letter}1",
                        f"Column {column_letter} is hidden.",
                    )
                )

        for row in worksheet.iter_rows():
            for cell in row:
                value = cell.value
                citation = _citation(worksheet.title, cell.coordinate)
                if _is_formula(value):
                    formula_graph[citation] = _formula_references(
                        value, worksheet.title, sheet_names
                    )
                    if _EXTERNAL_BOOK.search(value):
                        external_formula_cells.add(citation)
                        findings.append(
                            _finding(
                                "external_link",
                                worksheet.title,
                                cell.coordinate,
                                f"Formula contains an external-workbook reference: {value}",
                                severity="high",
                            )
                        )
                    continue
                if not _is_hardcoded_value(value):
                    continue
                neighbors = _opposite_formula_neighbors(
                    worksheet, cell.row, cell.column
                )
                if neighbors:
                    value_type = "numeric" if isinstance(value, Number) else "constant"
                    findings.append(
                        _finding(
                            "hardcoded_in_formula_range",
                            worksheet.title,
                            cell.coordinate,
                            (
                                f"{value_type.capitalize()} value {value!r} sits between "
                                "formula cells and may be an overwritten formula."
                            ),
                            severity="high",
                            related_cells=neighbors,
                        )
                    )

    for component in _strongly_connected_components(formula_graph):
        is_cycle = len(component) > 1 or (
            len(component) == 1 and component[0] in formula_graph[component[0]]
        )
        if not is_cycle:
            continue
        cells = sorted(component)
        first = cells[0]
        # The citation helper expects an unquoted title.  Use the workbook cell
        # citation directly here because ``first`` is already normalized.
        findings.append(
            {
                "type": "circular_reference",
                "category": "circular_reference",
                "severity": "high",
                "cell_citation": first,
                "related_cells": cells,
                "detail": (
                    "Textual formula-reference cycle detected among: "
                    + ", ".join(cells)
                    + ". Confirm in Excel because openpyxl does not calculate formulas."
                ),
            }
        )

    # Some workbooks carry an external-link package even if its formula is no
    # longer exposed in a cell.  This is kept distinct from cell-level links;
    # A1 is an explicit worksheet anchor rather than a claim that A1 has the link.
    packaged_links = list(getattr(workbook, "_external_links", ()) or ())
    if packaged_links and not external_formula_cells:
        anchor_sheet = workbook.worksheets[0].title
        for index, _link in enumerate(packaged_links, start=1):
            findings.append(
                _finding(
                    "external_link_package",
                    anchor_sheet,
                    "A1",
                    (
                        f"Workbook external-link relationship {index} exists; the A1 "
                        "citation is a workbook anchor because no referring cell was exposed."
                    ),
                    severity="high",
                )
            )

    counts: dict[str, int] = {}
    for item in findings:
        counts[item["type"]] = counts.get(item["type"], 0) + 1
    return {
        "status": "FINDINGS" if findings else "NO_STRUCTURAL_FINDINGS",
        "xlsx_path": str(path.resolve()),
        "worksheets_reviewed": len(workbook.worksheets),
        "finding_count": len(findings),
        "finding_counts_by_type": counts,
        "findings": findings,
        "unrecognized_inputs": [],
        "method": (
            "Read-only openpyxl structural pass over hidden dimensions/sheets, "
            "formula adjacency, textual formula dependencies, and external references."
        ),
        "limitations": [
            "openpyxl does not calculate formulas; circular references are detected textually.",
            "A constant is flagged only when immediately bracketed by formulas horizontally or vertically.",
            "Macros, named-range indirection, Power Query, data connections, and formula correctness require separate review.",
        ],
    }


def audit_seller_model(xlsx_path: str | Path | None) -> dict[str, Any]:
    """Audit a real Excel file and contain failures at the tool boundary."""

    try:
        return _audit_seller_model(xlsx_path)
    except Exception as exc:
        message = str(exc.args[0]) if isinstance(exc, KeyError) and exc.args else str(exc)
        return {"error": message or exc.__class__.__name__}


__all__ = ["audit_seller_model"]
